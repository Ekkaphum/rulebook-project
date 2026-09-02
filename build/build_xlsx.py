# -*- coding: utf-8 -*-
import sys, os, json, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from data_kb import S, KB
from data_aux import PRAISED, CRITICISED, TEMPLATES, STYLE_RULES, CHECKLIST, METRICS

HDR = PatternFill("solid", fgColor="1F3864")
HDRF = Font(color="FFFFFF", bold=True, size=11)
TITLE = Font(bold=True, size=16, color="1F3864")
SUB = Font(italic=True, size=10, color="595959")
WRAP = Alignment(wrap_text=True, vertical="top")
TOPL = Alignment(vertical="top")
THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)

wb = Workbook()

def sheet(name, headers, rows, widths, freeze="A2", link_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers)+1):
        cell = ws.cell(1, c); cell.fill = HDR; cell.font = HDRF
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = WRAP; cell.border = THIN
    for lc in link_cols:
        for r in range(2, ws.max_row+1):
            cell = ws.cell(r, lc)
            v = cell.value
            if isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v; cell.font = Font(color="0563C1", underline="single", size=10)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws

# ---------- 00 README ----------
ws = wb.active; ws.title = "00_README"
today = datetime.date.today().isoformat()
lines = [
 ("Board Game Rulebook Research Database", TITLE),
 (f"Compiled {today} · 106 primary sources · 127 knowledge entries", SUB),
 ("", None),
 ("WHAT THIS IS", Font(bold=True, size=12)),
 ("A structured knowledge base on board game rulebooks: what makes them good, what makes them bad, how the community", None),
 ("criticises them, and what practitioners recommend. Every row carries a source title and a clickable URL so any claim", None),
 ("can be traced back to its origin.", None),
 ("", None),
 ("HOW THE RESEARCH WAS DONE", Font(bold=True, size=12)),
 ("· Targeted harvesting of BoardGameGeek forum threads, geeklists and blogs on rulebook design, criticism and process.", None),
 ("· Designer/publisher primary sources: Stonemaier Games, Resonym, Leder Games, Fantasy Flight, Gaming Rules! (Paul Grogan).", None),
 ("· Professional style guides: Michael 'Curby' Lee's Board Game Editing Style Guide (~13,000 words, CC-BY-NC-SA),", None),
 ("  the Stonemaier Games Style Guide, and Kathleen Mercury's rules template.", None),
 ("· Trade and review press: Meeple Mountain, Punchboard, Going Analog, I Slay the Dragon, Board Game Design Lab.", None),
 ("· Accessibility guidance: WCAG-derived contrast targets, dyslexia typography guidance, Stonemaier accessibility model.", None),
 ("· BGG pages were retrieved through a text-extraction proxy because BoardGameGeek serves a Cloudflare challenge to", None),
 ("  automated clients, and the BGG XML API now returns 401 Unauthorized for anonymous forum/geeklist requests.", None),
 ("", None),
 ("SCOPE AND LIMITS — please read", Font(bold=True, size=12, color="C00000")),
 ("· This is TARGETED research, not an exhaustive crawl. Reading every BGG forum post from the last five years is not", None),
 ("  possible: BGG hosts millions of posts, has no bulk export, and rate-limits automated access.", None),
 ("· Threads were selected by relevance to rulebooks (quality, criticism, writing craft, structure, editing, testing,", None),
 ("  accessibility, localisation), spanning roughly 2007-2026 with emphasis on 2019-2026.", None),
 ("· Reddit could not be included: reddit.com blocks the crawler used here and is blocked by browsing policy.", None),
 ("· Community opinion is opinion. Rows in the 'Praised' and 'Criticised' sheets record what named sources said,", None),
 ("  not a verdict. Several games appear on both lists.", None),
 ("· Quotations are short excerpts for attribution; follow the URL for full context.", None),
 ("", None),
 ("SHEET GUIDE", Font(bold=True, size=12)),
 ("01_Sources           Every document harvested, with URL, era and size.", None),
 ("02_Knowledge_Base    127 findings organised into 14 categories, each with rationale, evidence and source.", None),
 ("03_Praised_Rulebooks 35 rulebooks named as exemplary, and exactly what was praised.", None),
 ("04_Criticised        33 rulebooks named as problematic, and the specific failure mode.", None),
 ("05_Structure_Templates  10 section-order templates from different authorities, side by side.", None),
 ("06_Style_Rules       26 concrete writing rules with a right/wrong example for each.", None),
 ("07_Checklist         46-item pre-print QA checklist across 8 phases.", None),
 ("08_Metrics           16 quantitative reference points (font sizes, contrast, rates, timings).", None),
 ("09_Failure_Modes     The failure taxonomy pulled out for quick scanning.", None),
 ("", None),
 ("LICENCE NOTE", Font(bold=True, size=12)),
 ("Michael Lee's style guide is CC-BY-NC-SA; material derived from it is attributed here and in the companion documents.", None),
 ("All other material is summarised and attributed, not reproduced.", None),
]
for i,(txt,f) in enumerate(lines, start=1):
    c = ws.cell(i,1,txt)
    if f: c.font = f
ws.column_dimensions['A'].width = 125

# ---------- 01 Sources ----------
auto = json.load(open(os.path.join(os.path.dirname(__file__),'sources_auto.json')))
def platform(u):
    if 'boardgamegeek.com/geeklist' in u: return 'BGG GeekList'
    if 'boardgamegeek.com/blog' in u: return 'BGG Blog'
    if 'boardgamegeek.com/wiki' in u: return 'BGG Wiki'
    if 'boardgamegeek.com' in u: return 'BGG Forum Thread'
    if 'stonemaiergames' in u: return 'Publisher Blog'
    if 'docs.google' in u or 'cur.by' in u: return 'Style Guide'
    if 'substack' in u: return 'Newsletter'
    return 'Web / Blog'
srows = []
for i, r in enumerate(sorted(auto, key=lambda x: x['slug']), start=1):
    srows.append((f"SRC-{i:03d}", platform(r['url']), r['title'], r['url'], r['years'], r['chars']))
srows.append((f"SRC-{len(srows)+1:03d}", 'Style Guide', 'Board Game Editing Style Guide (Michael "Curby" Lee) — full text export',
              'https://cur.by/styleguide', '2020-2025', 84176))
sheet("01_Sources", ["ID","Type","Title","URL","Years seen in thread","Chars harvested"],
      srows, [10,18,64,62,20,16], link_cols=(4,))

# ---------- 02 Knowledge base ----------
krows = []
for tid, cat, topic, principle, detail, evidence, sk in KB:
    st, su = S.get(sk, (sk, ''))
    krows.append((tid, cat, topic, principle, detail, evidence, st, su))
sheet("02_Knowledge_Base",
      ["ID","Category","Topic","Principle","Detail & rationale","Evidence / quote","Source","Source URL"],
      krows, [8,20,34,52,66,58,44,54], link_cols=(8,))

# ---------- 03 / 04 ----------
prows = [(g, a, w, S.get(k,(k,''))[0], S.get(k,(k,''))[1]) for g,a,w,k in PRAISED]
sheet("03_Praised_Rulebooks", ["Game / Publisher","Praised for","What specifically","Source","Source URL"],
      prows, [40,28,80,42,54], link_cols=(5,))
crows = [(g, a, w, S.get(k,(k,''))[0], S.get(k,(k,''))[1]) for g,a,w,k in CRITICISED]
sheet("04_Criticised_Rulebooks", ["Game / Publisher","Failure mode","What specifically","Source","Source URL"],
      crows, [40,30,84,42,54], link_cols=(5,))

# ---------- 05 templates ----------
maxlen = max(len(v) for v in TEMPLATES.values())
names = list(TEMPLATES.keys())
trows = []
for i in range(maxlen):
    trows.append([i+1] + [ (TEMPLATES[n][i] if i < len(TEMPLATES[n]) else "") for n in names ])
sheet("05_Structure_Templates", ["#"]+names, trows, [4]+[30]*len(names), freeze="B2")

# ---------- 06 style rules ----------
srows2 = [(a, r, d, dn, S.get(k,(k,''))[0], S.get(k,(k,''))[1]) for a,r,d,dn,k in STYLE_RULES]
sheet("06_Style_Rules", ["Area","Rule","Do (example)","Don't (example)","Source","Source URL"],
      srows2, [18,50,50,50,40,52], link_cols=(6,))

# ---------- 07 checklist ----------
chrows = [(p, i, w, "") for p,i,w in CHECKLIST]
ws = sheet("07_Checklist", ["Phase","Check","Why it matters","Done? (Y/N)"], chrows, [26,72,64,14])

# ---------- 08 metrics ----------
mrows = [(m, v, c, S.get(k,(k,''))[0], S.get(k,(k,''))[1]) for m,v,c,k in METRICS]
sheet("08_Metrics", ["Metric","Value","Context","Source","Source URL"], mrows, [40,44,50,42,54], link_cols=(5,))

# ---------- 09 failure modes ----------
frows = []
for tid, cat, topic, principle, detail, evidence, sk in KB:
    if cat == "Failure Modes":
        st, su = S.get(sk, (sk,''))
        frows.append((tid, topic, principle, detail, evidence, st, su))
sheet("09_Failure_Modes", ["ID","Failure mode","Description","Detail & examples","Evidence / quote","Source","Source URL"],
      frows, [8,40,56,72,58,42,52], link_cols=(7,))

out = "deliverables/Rulebook_Research_Database.xlsx"
wb.save(out)
print("saved", out, os.path.getsize(out), "bytes")
print("sheets:", wb.sheetnames)
