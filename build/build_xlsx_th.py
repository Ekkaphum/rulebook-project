# -*- coding: utf-8 -*-
import sys, os, json, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from data_kb import S
from data_kb_th import KB_TH
from data_aux_th import PRAISED_TH, CRITICISED_TH, TEMPLATES_TH, STYLE_RULES_TH, CHECKLIST_TH, METRICS_TH

FONT = "Tahoma"
HDR = PatternFill("solid", fgColor="1F3864")
HDRF = Font(name=FONT, color="FFFFFF", bold=True, size=10.5)
TITLE = Font(name=FONT, bold=True, size=16, color="1F3864")
SUB = Font(name=FONT, italic=True, size=10, color="595959")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)
BODY = Font(name=FONT, size=10)

wb = Workbook()

def sheet(name, headers, rows, widths, link_cols=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in range(1, len(headers)+1):
        cell = ws.cell(1, c); cell.fill = HDR; cell.font = HDRF
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(list(r))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment = WRAP; cell.border = THIN; cell.font = BODY
    for lc in link_cols:
        for r in range(2, ws.max_row+1):
            cell = ws.cell(r, lc); v = cell.value
            if isinstance(v, str) and v.startswith("http"):
                cell.hyperlink = v
                cell.font = Font(name=FONT, color="0563C1", underline="single", size=9.5)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws

# ---------- 00 อ่านก่อน ----------
ws = wb.active; ws.title = "00_อ่านก่อน"
today = datetime.date.today().isoformat()
lines = [
 ("ฐานข้อมูลงานวิจัย Rulebook บอร์ดเกม", TITLE),
 (f"เรียบเรียง {today} · 106 แหล่งข้อมูล · 127 องค์ความรู้", SUB),
 ("", None),
 ("นี่คืออะไร", Font(name=FONT, bold=True, size=12)),
 ("ฐานข้อมูลเชิงโครงสร้างเรื่อง rulebook บอร์ดเกม: อะไรทำให้ดี อะไรทำให้แย่ ชุมชนวิจารณ์มันยังไง และคนทำงานจริงแนะนำอะไร", None),
 ("ทุกแถวมีชื่อแหล่งที่มาและ URL ที่คลิกได้ เพื่อให้ตามกลับไปตรวจสอบต้นทางได้ทุกข้อความ", None),
 ("", None),
 ("งานวิจัยนี้ทำยังไง", Font(name=FONT, bold=True, size=12)),
 ("· เจาะเธรด geeklist และบล็อกบน BoardGameGeek ที่เกี่ยวกับการออกแบบ วิจารณ์ และกระบวนการทำ rulebook", None),
 ("· แหล่งปฐมภูมิจากนักออกแบบและสำนักพิมพ์: Stonemaier Games, Resonym, Leder Games, Fantasy Flight, Gaming Rules! (Paul Grogan)", None),
 ("· Style guide อาชีพ: Board Game Editing Style Guide ของ Michael 'Curby' Lee (~13,000 คำ, CC-BY-NC-SA),", None),
 ("  Stonemaier Games Style Guide และ template ของ Kathleen Mercury", None),
 ("· สื่อและนักรีวิว: Meeple Mountain, Punchboard, Going Analog, I Slay the Dragon, Board Game Design Lab", None),
 ("· แนวทางด้านการเข้าถึง: เป้าหมายความต่างสีที่มาจาก WCAG, แนวทางตัวอักษรสำหรับ dyslexia, โมเดลการเข้าถึงของ Stonemaier", None),
 ("· หน้าเว็บ BGG ถูกดึงผ่าน text-extraction proxy เพราะ BoardGameGeek ส่ง Cloudflare challenge ให้ client อัตโนมัติ", None),
 ("  และ BGG XML API คืนค่า 401 Unauthorized สำหรับคำขอ forum/geeklist แบบไม่ล็อกอินแล้ว", None),
 ("", None),
 ("ขอบเขตและข้อจำกัด — กรุณาอ่าน", Font(name=FONT, bold=True, size=12, color="C00000")),
 ("· นี่คืองานวิจัยแบบเจาะจง ไม่ใช่การไล่เก็บทั้งหมด การอ่านทุกโพสต์บน BGG ย้อนหลังห้าปีเป็นไปไม่ได้:", None),
 ("  BGG มีโพสต์หลักล้าน ไม่มีช่องทาง export ทั้งชุด และจำกัดการเข้าถึงแบบอัตโนมัติ", None),
 ("· เธรดถูกคัดเลือกตามความเกี่ยวข้องกับ rulebook (คุณภาพ การวิจารณ์ งานเขียน โครงสร้าง การแก้ไข การทดสอบ", None),
 ("  การเข้าถึง และการแปล) ครอบคลุมราวปี 2007-2026 โดยเน้นหนักที่ 2019-2026", None),
 ("· Reddit รวมไม่ได้: reddit.com บล็อกทั้ง crawler ที่ใช้และถูกบล็อกโดยนโยบายการเข้าเว็บ", None),
 ("· ความเห็นในชุมชนคือความเห็น แถวในชีต 'ชื่นชม' และ 'ถูกวิจารณ์' บันทึกสิ่งที่แหล่งที่ระบุชื่อพูดไว้", None),
 ("  ไม่ใช่คำตัดสิน มีหลายเกมที่ปรากฏในทั้งสองรายการ", None),
 ("· คำพูดที่ยกมาเป็นข้อความสั้นเพื่อการอ้างอิง กรุณาตาม URL ไปอ่านบริบทเต็ม", None),
 ("", None),
 ("คู่มือชีต", Font(name=FONT, bold=True, size=12)),
 ("01_แหล่งข้อมูล        เอกสารทุกชิ้นที่เก็บมา พร้อม URL ช่วงปี และขนาด", None),
 ("02_องค์ความรู้        127 ข้อค้นพบใน 14 หมวด แต่ละข้อมีเหตุผล หลักฐาน และแหล่งที่มา", None),
 ("03_rulebook_ที่ควรศึกษา   35 เล่มที่ถูกยกเป็นตัวอย่างดี และสิ่งที่ถูกชมอย่างเจาะจง", None),
 ("04_rulebook_ที่ถูกวิจารณ์  33 เล่มที่ถูกระบุว่ามีปัญหา และรูปแบบความล้มเหลวที่เฉพาะเจาะจง", None),
 ("05_template_โครงสร้าง  10 ลำดับหัวข้อจากแหล่งอ้างอิงต่างกัน วางเทียบกัน", None),
 ("06_กฎการเขียน        26 กฎการเขียนที่เป็นรูปธรรม พร้อมตัวอย่างถูก/ผิด", None),
 ("07_เช็กลิสต์          เช็กลิสต์ QA ก่อนพิมพ์ 46 ข้อ ใน 8 ระยะ", None),
 ("08_ตัวเลขอ้างอิง       16 จุดอ้างอิงเชิงปริมาณ (ขนาดตัวอักษร ความต่างสี ค่าจ้าง เวลา)", None),
 ("09_รูปแบบความล้มเหลว   ระบบจำแนกความล้มเหลว แยกออกมาให้กวาดตาได้เร็ว", None),
 ("", None),
 ("หมายเหตุเรื่องสัญญาอนุญาต", Font(name=FONT, bold=True, size=12)),
 ("style guide ของ Michael Lee ใช้สัญญาอนุญาต CC-BY-NC-SA เนื้อหาที่ต่อยอดจากมันถูกอ้างอิงไว้ทั้งในไฟล์นี้และในเอกสารประกอบ", None),
 ("แหล่งอื่นทั้งหมดถูกสรุปและอ้างอิง ไม่ได้คัดลอกมาทั้งชิ้น", None),
]
for i,(txt,f) in enumerate(lines, start=1):
    c = ws.cell(i,1,txt); c.font = f if f else Font(name=FONT, size=10.5)
ws.column_dimensions['A'].width = 118

# ---------- 01 แหล่งข้อมูล ----------
auto = json.load(open(os.path.join(os.path.dirname(__file__),'sources_auto.json')))
def plat(u):
    if 'boardgamegeek.com/geeklist' in u: return 'BGG GeekList'
    if 'boardgamegeek.com/blog' in u: return 'BGG บล็อก'
    if 'boardgamegeek.com/wiki' in u: return 'BGG วิกิ'
    if 'boardgamegeek.com' in u: return 'BGG เธรดฟอรั่ม'
    if 'stonemaiergames' in u: return 'บล็อกสำนักพิมพ์'
    if 'docs.google' in u or 'cur.by' in u: return 'Style Guide'
    if 'substack' in u: return 'จดหมายข่าว'
    return 'เว็บ / บล็อก'
srows = [(f"SRC-{i:03d}", plat(r['url']), r['title'], r['url'], r['years'], r['chars'])
         for i, r in enumerate(sorted(auto, key=lambda x: x['slug']), start=1)]
srows.append((f"SRC-{len(srows)+1:03d}", 'Style Guide',
              'Board Game Editing Style Guide (Michael "Curby" Lee) — ข้อความเต็ม',
              'https://cur.by/styleguide', '2020-2025', 84176))
sheet("01_แหล่งข้อมูล", ["รหัส","ประเภท","ชื่อเอกสาร (ต้นฉบับภาษาอังกฤษ)","URL","ปีที่พบในเธรด","จำนวนอักขระที่เก็บ"],
      srows, [10,18,64,62,20,18], link_cols=(4,))

# ---------- 02 องค์ความรู้ ----------
krows = [(a,b,c,d,e,f, S.get(g,(g,''))[0], S.get(g,(g,''))[1]) for a,b,c,d,e,f,g in KB_TH]
sheet("02_องค์ความรู้", ["รหัส","หมวด","หัวข้อ","หลักการ","รายละเอียดและเหตุผล","หลักฐาน / คำพูด","แหล่งที่มา","URL"],
      krows, [8,20,36,54,68,58,44,54], link_cols=(8,))

# ---------- 03 / 04 ----------
sheet("03_rulebook_ที่ควรศึกษา", ["เกม / สำนักพิมพ์","ถูกชมเรื่อง","รายละเอียด","แหล่งที่มา","URL"],
      [(g,a,d,S.get(k,(k,''))[0],S.get(k,(k,''))[1]) for g,a,d,k in PRAISED_TH],
      [40,28,80,42,54], link_cols=(5,))
sheet("04_rulebook_ที่ถูกวิจารณ์", ["เกม / สำนักพิมพ์","รูปแบบความล้มเหลว","รายละเอียด","แหล่งที่มา","URL"],
      [(g,a,d,S.get(k,(k,''))[0],S.get(k,(k,''))[1]) for g,a,d,k in CRITICISED_TH],
      [40,30,84,42,54], link_cols=(5,))

# ---------- 05 templates ----------
names = list(TEMPLATES_TH.keys())
maxlen = max(len(v) for v in TEMPLATES_TH.values())
trows = [[i+1] + [(TEMPLATES_TH[n][i] if i < len(TEMPLATES_TH[n]) else "") for n in names] for i in range(maxlen)]
sheet("05_template_โครงสร้าง", ["ลำดับ"]+names, trows, [6]+[30]*len(names))

# ---------- 06 style rules ----------
sheet("06_กฎการเขียน", ["ด้าน","กฎ","ควรทำ (ตัวอย่าง)","ไม่ควรทำ (ตัวอย่าง)","แหล่งที่มา","URL"],
      [(a,r,d,dn,S.get(k,(k,''))[0],S.get(k,(k,''))[1]) for a,r,d,dn,k in STYLE_RULES_TH],
      [20,50,50,50,40,52], link_cols=(6,))

# ---------- 07 checklist ----------
sheet("07_เช็กลิสต์", ["ระยะ","สิ่งที่ต้องตรวจ","ทำไมถึงสำคัญ","เสร็จแล้ว? (Y/N)"],
      [(p,i,w,"") for p,i,w in CHECKLIST_TH], [28,74,62,16])

# ---------- 08 metrics ----------
sheet("08_ตัวเลขอ้างอิง", ["ตัวชี้วัด","ค่า","บริบท","แหล่งที่มา","URL"],
      [(m,v,c,S.get(k,(k,''))[0],S.get(k,(k,''))[1]) for m,v,c,k in METRICS_TH],
      [42,46,50,42,54], link_cols=(5,))

# ---------- 09 failure modes ----------
frows = [(a,c,d,e,f,S.get(g,(g,''))[0],S.get(g,(g,''))[1]) for a,b,c,d,e,f,g in KB_TH if b=="รูปแบบความล้มเหลว"]
sheet("09_รูปแบบความล้มเหลว", ["รหัส","รูปแบบ","คำอธิบาย","รายละเอียดและตัวอย่าง","หลักฐาน / คำพูด","แหล่งที่มา","URL"],
      frows, [8,42,56,74,58,42,52], link_cols=(7,))

out = "deliverables/ฐานข้อมูลงานวิจัย Rulebook.xlsx"
wb.save(out)
print("saved", out, os.path.getsize(out), "bytes")
print("sheets:", wb.sheetnames)
